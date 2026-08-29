"""
عرض دفتر العميل
Customer ledger view
"""
from PySide6.QtWidgets import QWidget, QVBoxLayout, QHBoxLayout, QTableWidget, QTableWidgetItem, QPushButton, QLabel, QGroupBox, QHeaderView, QMessageBox, QMenu
from PySide6.QtCore import Qt, Signal
from PySide6.QtGui import QFont, QColor, QAction
from decimal import Decimal
from datetime import datetime
from typing import List
import csv
import json

from core.models import Customer, Entry, Summary
from core.storage import StorageManager
from core.summaries import calculate_customer_summary, calculate_running_balances
from core.utils import format_currency, format_datetime_arabic
from .dialogs import AddEntryDialog, ExportDialog, EditEntryDialog
from .table_items import DateTimeItem, NumericItem, sorted_table


class CustomerLedgerView(QWidget):
    """
    عرض دفتر العميل مع الإدخالات والملخصات
    Customer ledger view with entries and summaries
    """
    
    # إشارة للعودة إلى قائمة العملاء
    back_requested = Signal()
    
    def __init__(self, storage: StorageManager, parent=None):
        super().__init__(parent)
        self.storage = storage
        self.customer: Customer = None
        self.entries: List[Entry] = []
        self.summary: Summary = None
        self.current_filter = "all"  # all, today, month
        self.setup_ui()
    
    def setup_ui(self):
        """إعداد الواجهة"""
        layout = QVBoxLayout()
        
        # الشريط العلوي مع اسم العميل
        top_bar = QHBoxLayout()
        
        self.customer_name_label = QLabel("")
        title_font = QFont()
        title_font.setPointSize(16)
        title_font.setBold(True)
        self.customer_name_label.setFont(title_font)
        top_bar.addWidget(self.customer_name_label)
        
        top_bar.addStretch()
        
        # زر العودة
        self.back_btn = QPushButton("← العودة")
        self.back_btn.clicked.connect(self.back_requested.emit)
        top_bar.addWidget(self.back_btn)
        
        layout.addLayout(top_bar)
        
        # بطاقات الملخص
        summary_layout = QHBoxLayout()
        
        # الرصيد الحالي
        self.balance_box = self.create_summary_card("الرصيد الحالي", "0.00 ج.م")
        summary_layout.addWidget(self.balance_box)
        
        # مبيعات اليوم
        self.today_sales_box = self.create_summary_card("مبيعات اليوم", "0.00 ج.م")
        summary_layout.addWidget(self.today_sales_box)
        
        # مبيعات الشهر
        self.month_sales_box = self.create_summary_card("مبيعات الشهر", "0.00 ج.م")
        summary_layout.addWidget(self.month_sales_box)
        
        layout.addLayout(summary_layout)
        
        # أزرار التصفية
        filter_layout = QHBoxLayout()
        
        filter_label = QLabel("عرض:")
        filter_layout.addWidget(filter_label)
        
        self.all_btn = QPushButton("الكل")
        self.all_btn.setCheckable(True)
        self.all_btn.setChecked(True)
        self.all_btn.clicked.connect(lambda: self.set_filter("all"))
        filter_layout.addWidget(self.all_btn)
        
        self.today_btn = QPushButton("اليوم")
        self.today_btn.setCheckable(True)
        self.today_btn.clicked.connect(lambda: self.set_filter("today"))
        filter_layout.addWidget(self.today_btn)
        
        self.month_btn = QPushButton("هذا الشهر")
        self.month_btn.setCheckable(True)
        self.month_btn.clicked.connect(lambda: self.set_filter("month"))
        filter_layout.addWidget(self.month_btn)
        
        filter_layout.addStretch()
        
        layout.addLayout(filter_layout)
        
        # جدول الإدخالات
        self.table = QTableWidget()
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels([
            "التاريخ والوقت",
            "النوع",
            "الوصف",
            "طريقة الدفع",
            "المبلغ",
            "الرصيد بعد العملية"
        ])
        
        # إعدادات الجدول
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        
        # تمديد الأعمدة
        header = self.table.horizontalHeader()
        header.setStretchLastSection(True)
        for i in range(6):
            header.setSectionResizeMode(i, QHeaderView.Stretch)
        
        # تفعيل قائمة السياق (النقر بالزر الأيمن)
        self.table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self.show_entry_context_menu)
        
        layout.addWidget(self.table)
        
        # أزرار الإجراءات
        actions_layout = QHBoxLayout()
        
        self.add_purchase_btn = QPushButton("+ إضافة دين")
        self.add_purchase_btn.clicked.connect(lambda: self.add_entry("purchase"))
        actions_layout.addWidget(self.add_purchase_btn)
        
        self.add_payment_btn = QPushButton("+ إضافة دفع")
        self.add_payment_btn.clicked.connect(lambda: self.add_entry("payment"))
        actions_layout.addWidget(self.add_payment_btn)
        
        actions_layout.addStretch()
        
        self.export_btn = QPushButton("تصدير")
        self.export_btn.clicked.connect(self.export_ledger)
        actions_layout.addWidget(self.export_btn)
        
        layout.addLayout(actions_layout)
        
        self.setLayout(layout)
    
    def create_summary_card(self, title: str, value: str) -> QGroupBox:
        """إنشاء بطاقة ملخص"""
        box = QGroupBox(title)
        layout = QVBoxLayout()
        
        value_label = QLabel(value)
        value_font = QFont()
        value_font.setPointSize(18)
        value_font.setBold(True)
        value_label.setFont(value_font)
        value_label.setAlignment(Qt.AlignCenter)
        
        layout.addWidget(value_label)
        box.setLayout(layout)
        
        # حفظ مرجع إلى التسمية للتحديث لاحقاً
        box.value_label = value_label
        
        return box
    
    def set_customer(self, customer_id: str):
        """تعيين العميل الحالي"""
        # تحميل العميل
        self.customer = self.storage.get_customer(customer_id)
        if not self.customer:
            QMessageBox.warning(self, "خطأ", f"العميل غير موجود: {customer_id}")
            return
        
        # تحديث العنوان
        self.customer_name_label.setText(f"دفتر العميل: {self.customer.name}")
        
        # تحميل البيانات
        self.load_data()
    
    def load_data(self):
        """تحميل بيانات العميل"""
        if not self.customer:
            return
        
        # تحميل الإدخالات
        all_entries = self.storage.get_customer_entries(self.customer.id)
        
        # ترتيب حسب التاريخ (الأحدث أولاً)
        all_entries.sort(key=lambda e: e.datetime, reverse=True)
        
        self.entries = all_entries
        
        # حساب الملخص
        all_entries_for_summary = self.storage.load_entries()
        self.summary = calculate_customer_summary(self.customer.id, all_entries_for_summary)
        
        # تحديث الملخصات في التخزين
        summaries = self.storage.load_summaries()
        summaries[self.customer.id] = self.summary
        self.storage.save_summaries(summaries)
        
        # تحديث الواجهة
        self.update_summary_cards()
        self.update_table()
    
    def update_summary_cards(self):
        """تحديث بطاقات الملخص"""
        if not self.summary:
            return
        
        # الرصيد الحالي
        balance = self.summary.balance_decimal
        self.balance_box.value_label.setText(format_currency(balance))
        
        # تلوين الرصيد
        if balance > 0:
            self.balance_box.value_label.setStyleSheet("color: #dc3545;")  # أحمر
        elif balance < 0:
            self.balance_box.value_label.setStyleSheet("color: #28a745;")  # أخضر
        else:
            self.balance_box.value_label.setStyleSheet("color: #6c757d;")  # رمادي
        
        # مبيعات اليوم
        today_sales = self.summary.today_sales_decimal
        self.today_sales_box.value_label.setText(format_currency(today_sales))
        
        # مبيعات الشهر
        month_sales = self.summary.month_sales_decimal
        self.month_sales_box.value_label.setText(format_currency(month_sales))
    
    def set_filter(self, filter_type: str):
        """تعيين التصفية"""
        self.current_filter = filter_type
        
        # تحديث الأزرار
        self.all_btn.setChecked(filter_type == "all")
        self.today_btn.setChecked(filter_type == "today")
        self.month_btn.setChecked(filter_type == "month")
        
        # تحديث الجدول
        self.update_table()
    
    def update_table(self):
        """تحديث جدول الإدخالات"""
        # تصفية الإدخالات
        now = datetime.now()
        
        if self.current_filter == "today":
            filtered_entries = [
                e for e in self.entries 
                if e.datetime.date() == now.date()
            ]
        elif self.current_filter == "month":
            filtered_entries = [
                e for e in self.entries 
                if e.datetime.year == now.year and e.datetime.month == now.month
            ]
        else:
            filtered_entries = self.entries
        
        # حساب الأرصدة التراكمية من كل الحركات (لا المصفاة)، لأن الرصيد
        # التراكمي لا معنى له إلا بترتيب زمني كامل
        entries_sorted_asc = sorted(self.entries, key=lambda e: e.datetime)
        running_balances = calculate_running_balances(entries_sorted_asc)
        balance_map = {
            entry.id: balance
            for entry, balance in zip(entries_sorted_asc, running_balances)
        }

        # الترتيب معطَّل أثناء التعبئة، ثم يُعاد ويُطبَّق الافتراضي:
        # التاريخ تنازلياً (الأحدث أولاً) — الترتيب الطبيعي للدفتر
        with sorted_table(self.table, column=0, order=Qt.DescendingOrder):
            self.table.setRowCount(len(filtered_entries))

            for row, entry in enumerate(filtered_entries):
                # التاريخ والوقت — ترتيب زمني لا نصّي
                date_item = DateTimeItem(
                    format_datetime_arabic(entry.datetime), entry.datetime
                )
                # معرّف الحركة على العمود الأول: بعد الترتيب لا يوافق رقم
                # الصف موضع العنصر في القائمة، فالبحث يجب أن يكون بالمعرّف
                date_item.setData(Qt.UserRole, entry.id)
                self.table.setItem(row, 0, date_item)

                # النوع
                type_text = "دين" if entry.type == "purchase" else "دفع"
                type_item = QTableWidgetItem(type_text)
                if entry.type == "purchase":
                    type_item.setForeground(QColor(220, 53, 69))  # أحمر
                else:
                    type_item.setForeground(QColor(40, 167, 69))  # أخضر
                self.table.setItem(row, 1, type_item)

                # الوصف
                self.table.setItem(row, 2, QTableWidgetItem(entry.desc))

                # طريقة الدفع
                payment_text = "نقداً" if entry.pay_method == "cash" else "تحويل"
                self.table.setItem(row, 3, QTableWidgetItem(payment_text))

                # المبلغ — ترتيب رقمي
                self.table.setItem(row, 4, NumericItem(
                    format_currency(entry.amount_decimal), entry.amount_decimal
                ))

                # الرصيد بعد العملية — يُرتَّب زمنياً لا بقيمته، لأنه رصيد
                # تراكمي لا معنى له إلا بالترتيب الزمني
                balance = balance_map.get(entry.id, Decimal("0.00"))
                self.table.setItem(row, 5, DateTimeItem(
                    format_currency(balance), entry.datetime
                ))
    
    def add_entry(self, entry_type: str = "purchase"):
        """إضافة إدخال جديد"""
        if not self.customer:
            return
        
        # الحصول على الرصيد الحالي
        current_balance = self.summary.balance_decimal if self.summary else Decimal("0")
        
        # فتح نافذة الإضافة مع تمرير معلومات العميل والرصيد
        dialog = AddEntryDialog(
            self.customer.id, 
            self.customer.name,
            customer=self.customer,
            current_balance=current_balance,
            parent=self
        )
        
        # تعيين النوع الافتراضي
        if entry_type == "payment":
            dialog.type_combo.setCurrentIndex(1)
            dialog.on_type_changed(1)
        
        if dialog.exec():
            try:
                # حفظ الإدخال
                self.storage.add_entry(dialog.entry)
                
                # إعادة تحميل البيانات
                self.load_data()
                
                QMessageBox.information(
                    self, "نجح", 
                    f"تم إضافة {'الدين' if entry_type == 'purchase' else 'الدفع'} بنجاح"
                )
            except Exception as e:
                QMessageBox.warning(self, "خطأ", f"فشل في حفظ الإدخال: {str(e)}")
    
    def export_ledger(self):
        """تصدير دفتر العميل"""
        if not self.customer:
            return
        
        dialog = ExportDialog(self)
        if dialog.exec():
            try:
                if dialog.export_format == "csv":
                    self.export_to_csv(dialog.export_path)
                elif dialog.export_format == "json":
                    self.export_to_json(dialog.export_path)
                elif dialog.export_format == "pdf":
                    self.export_to_pdf(dialog.export_path)
                elif dialog.export_format == "excel":
                    self.export_to_excel(dialog.export_path)
                
                QMessageBox.information(
                    self, "نجح", 
                    f"تم تصدير الدفتر بنجاح إلى:\n{dialog.export_path}"
                )
            except Exception as e:
                QMessageBox.warning(self, "خطأ", f"فشل في التصدير: {str(e)}")
    
    def export_to_csv(self, path: str):
        """تصدير إلى CSV"""
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            
            # الرؤوس
            writer.writerow([
                "التاريخ والوقت", "النوع", "الوصف", 
                "طريقة الدفع", "المبلغ"
            ])
            
            # البيانات
            for entry in self.entries:
                writer.writerow([
                    format_datetime_arabic(entry.datetime),
                    "دين" if entry.type == "purchase" else "دفع",
                    entry.desc,
                    "نقداً" if entry.pay_method == "cash" else "تحويل",
                    str(entry.amount_decimal)
                ])
    
    def export_to_json(self, path: str):
        """تصدير إلى JSON"""
        data = {
            "customer": {
                "id": self.customer.id,
                "name": self.customer.name
            },
            "summary": self.summary.to_dict(),
            "entries": [e.to_dict() for e in self.entries]
        }
        
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    
    def export_to_pdf(self, path: str):
        """تصدير إلى PDF"""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib import colors
            from reportlab.lib.units import cm
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            import os
            
            # دالة مساعدة لتنسيق العملة في PDF
            def format_currency_pdf(amount):
                return f"{amount:,.2f} ILS"
            
            # دالة لإعادة تشكيل النص العربي لربط الأحرف
            def fix_arabic_text(text):
                """إعادة تشكيل النص العربي لربط الأحرف بشكل صحيح"""
                if not text:
                    return text
                try:
                    # محاولة استخدام arabic_reshaper إذا كان متوفراً
                    import arabic_reshaper
                    from bidi.algorithm import get_display
                    
                    # إعادة تشكيل النص العربي لربط الأحرف
                    reshaped_text = arabic_reshaper.reshape(text)
                    # عكس النص للعرض من اليسار لليمين
                    bidi_text = get_display(reshaped_text)
                    return bidi_text
                except ImportError:
                    # في حالة عدم توفر المكتبات، استخدم عكس بسيط
                    return text[::-1]
            
            # محاولة تحميل خط عربي من النظام
            arabic_font_registered = False
            try:
                font_paths = [
                    r"C:\Windows\Fonts\arial.ttf",
                    r"C:\Windows\Fonts\tahoma.ttf",
                    r"C:\Windows\Fonts\arialuni.ttf",
                ]
                
                for font_path in font_paths:
                    if os.path.exists(font_path):
                        pdfmetrics.registerFont(TTFont('Arabic', font_path))
                        arabic_font_registered = True
                        break
            except:
                pass
            
            # إنشاء ملف PDF
            doc = SimpleDocTemplate(path, pagesize=A4)
            elements = []
            
            # الأنماط
            styles = getSampleStyleSheet()
            
            # عنوان الدفتر
            title = Paragraph("<b>Customer Ledger Report</b>", styles['Title'])
            elements.append(title)
            elements.append(Spacer(1, 0.3*cm))
            
            # معلومات العميل
            customer_name_fixed = fix_arabic_text(self.customer.name) if arabic_font_registered else self.customer.name
            customer_id_fixed = fix_arabic_text(self.customer.id) if arabic_font_registered else self.customer.id
            
            if arabic_font_registered:
                customer_info = Paragraph(
                    f"<b>Customer Name:</b> <font name='Arabic'>{customer_name_fixed}</font><br/>"
                    f"<b>Customer ID:</b> <font name='Arabic'>{customer_id_fixed}</font><br/>"
                    f"<b>Total Balance:</b> {format_currency_pdf(self.summary.balance_decimal)}",
                    styles['Normal']
                )
            else:
                customer_info = Paragraph(
                    f"<b>Customer Name:</b> {self.customer.name}<br/>"
                    f"<b>Customer ID:</b> {self.customer.id}<br/>"
                    f"<b>Total Balance:</b> {format_currency_pdf(self.summary.balance_decimal)}",
                    styles['Normal']
                )
            elements.append(customer_info)
            elements.append(Spacer(1, 0.5*cm))
            
            # البيانات
            data = [
                ["Date/Time", "Type", "Description", "Amount"]
            ]
            
            for entry in self.entries:
                entry_type = "Debt" if entry.type == "purchase" else "Payment"
                desc_fixed = fix_arabic_text(entry.desc) if arabic_font_registered else entry.desc
                data.append([
                    format_datetime_arabic(entry.datetime),
                    entry_type,
                    desc_fixed,
                    format_currency_pdf(entry.amount_decimal)
                ])
            
            # إنشاء الجدول
            table = Table(data)
            
            if arabic_font_registered:
                table_style = TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTNAME', (0, 1), (2, -1), 'Arabic'),  # استخدام الخط العربي للوصف
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('FONTSIZE', (0, 1), (-1, -1), 11),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ])
            else:
                table_style = TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ])
            
            table.setStyle(table_style)
            elements.append(table)
            elements.append(Spacer(1, 1*cm))
            
            # تذييل مع التاريخ
            footer = Paragraph(
                f"<b>Generated:</b> {datetime.now().strftime('%Y/%m/%d %I:%M %p')}",
                styles['Normal']
            )
            elements.append(footer)
            
            # بناء PDF
            doc.build(elements)
            
        except ImportError:
            raise Exception("تصدير PDF يتطلب مكتبة reportlab. يرجى استخدام CSV أو JSON بدلاً من ذلك.")
    
    def export_to_excel(self, path: str):
        """تصدير إلى Excel"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter
            
            wb = Workbook()
            ws = wb.active
            ws.title = "دفتر العميل"
            
            # إعداد الأنماط
            header_font = Font(bold=True, color="FFFFFF", size=12)
            header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # معلومات العميل
            ws['A1'] = "اسم العميل:"
            ws['B1'] = self.customer.name
            ws['A2'] = "رقم العميل:"
            ws['B2'] = self.customer.id
            ws['A3'] = "الرصيد الحالي:"
            ws['B3'] = float(self.summary.balance_decimal)
            ws['A1'].font = Font(bold=True)
            ws['A2'].font = Font(bold=True)
            ws['A3'].font = Font(bold=True)
            
            # صف فارغ
            start_row = 5
            
            # الرؤوس
            headers = ["التاريخ والوقت", "النوع", "الوصف", "طريقة الدفع", "المبلغ", "الملاحظات"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=start_row, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            # البيانات
            for row_idx, entry in enumerate(self.entries, start_row + 1):
                ws.cell(row=row_idx, column=1, value=format_datetime_arabic(entry.datetime)).border = thin_border
                ws.cell(row=row_idx, column=2, value="دين" if entry.type == "purchase" else "دفع").border = thin_border
                ws.cell(row=row_idx, column=3, value=entry.desc).border = thin_border
                ws.cell(row=row_idx, column=4, value="نقداً" if entry.pay_method == "cash" else "تحويل").border = thin_border
                ws.cell(row=row_idx, column=5, value=float(entry.amount_decimal)).border = thin_border
                ws.cell(row=row_idx, column=6, value=entry.notes if entry.notes else "").border = thin_border
            
            # ضبط عرض الأعمدة
            for col in range(1, 7):
                ws.column_dimensions[get_column_letter(col)].width = 18
            
            # حفظ الملف
            wb.save(path)
            
        except ImportError:
            raise Exception("تصدير Excel يتطلب مكتبة openpyxl. يرجى استخدام CSV أو JSON بدلاً من ذلك.")
    
    def show_entry_context_menu(self, position):
        """
        عرض قائمة السياق للإدخال
        Show context menu for entry
        """
        # الحصول على الصف المحدد
        row = self.table.rowAt(position.y())
        if row < 0:
            return

        # البحث بالمعرّف لا برقم الصف: مع تفعيل الترتيب لم يعد رقم الصف
        # يوافق موضع الحركة في القائمة، فالفهرسة تفتح الحركة الخطأ
        item = self.table.item(row, 0)
        if item is None:
            return

        entry_id = item.data(Qt.UserRole)
        entry = next((e for e in self.entries if e.id == entry_id), None)
        if entry is None:
            return

        # إنشاء القائمة
        menu = QMenu(self)

        # تعديل الإدخال
        edit_action = QAction("✏️ تعديل", self)
        edit_action.triggered.connect(lambda: self.edit_entry(entry))
        menu.addAction(edit_action)
        
        menu.addSeparator()
        
        # حذف الإدخال
        delete_action = QAction("🗑️ حذف", self)
        delete_action.triggered.connect(lambda: self.delete_entry(entry))
        menu.addAction(delete_action)
        
        # عرض القائمة
        menu.exec(self.table.viewport().mapToGlobal(position))
    
    def edit_entry(self, entry: Entry, row: int = -1):
        """
        تعديل حركة موجودة
        Edit an existing entry
        """
        dialog = EditEntryDialog(entry, self.customer.name, self)

        if not dialog.exec():
            return

        try:
            updated = dialog.updated_entry
            # الحوار يبني Entry جديداً فيحصل على معرّف جديد تلقائياً؛
            # نعيد الهوية الأصلية وإلا فشل التحديث بـ KeyError
            updated.id = entry.id

            self.storage.update_entry(updated)
            self.load_data()
            QMessageBox.information(self, "نجح", "تم تعديل الحركة بنجاح")
        except KeyError:
            QMessageBox.warning(self, "خطأ", "لم يتم العثور على الحركة")
        except Exception as e:
            QMessageBox.warning(self, "خطأ", f"فشل في تعديل الحركة: {str(e)}")

    def delete_entry(self, entry: Entry):
        """
        حذف حركة
        Delete an entry
        """
        entry_type_ar = "الدين" if entry.type == "purchase" else "الدفعة"

        reply = QMessageBox.question(
            self, "تأكيد الحذف",
            f"هل أنت متأكد من حذف {entry_type_ar}؟\n\n"
            f"المبلغ: {format_currency(entry.amount_decimal)}\n"
            f"الوصف: {entry.desc}\n"
            f"التاريخ: {format_datetime_arabic(entry.datetime)}",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )

        if reply != QMessageBox.Yes:
            return

        try:
            self.storage.delete_entry(entry.id)
            self.load_data()
            QMessageBox.information(self, "نجح", f"تم حذف {entry_type_ar} بنجاح")
        except KeyError:
            QMessageBox.warning(self, "خطأ", "لم يتم العثور على الحركة")
        except Exception as e:
            QMessageBox.warning(self, "خطأ", f"فشل في حذف الحركة: {str(e)}")
